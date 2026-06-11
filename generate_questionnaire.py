# -*- coding: utf-8 -*-
"""Generate the bilingual (CN/EN) dissertation questionnaire as .docx,
following the structure of the IKEA example questionnaire."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ---------- base style ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)

GREY = RGBColor(0x7F, 0x7F, 0x7F)


def set_ea(run):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    set_ea(r)
    p.space_before = Pt(10)
    return p


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    set_ea(r)
    return p


def para(text, bold=False, italic=False, grey=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if grey:
        r.font.color.rgb = GREY
    if size:
        r.font.size = Pt(size)
    set_ea(r)
    return p


def note(text):
    """Admin/design note - grey italic, not shown to respondents."""
    return para(text, italic=True, grey=True, size=9.5)


def options(opts):
    for o in opts:
        para("○ " + o)


def checkboxes(opts):
    for o in opts:
        para("□ " + o)


def blank_line():
    para("____________________________________________")


def matrix(items, header="题项 / Item"):
    """7-point matrix table: item text + columns 1-7."""
    table = doc.add_table(rows=1 + len(items), cols=8)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = header
    for i in range(7):
        hdr[i + 1].text = str(i + 1)
    for c in table.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                set_ea(r)
    for ri, (code, cn, en) in enumerate(items, start=1):
        row = table.rows[ri].cells
        p = row[0].paragraphs[0]
        r1 = p.add_run(f"{code} {cn}")
        set_ea(r1)
        p2 = row[0].add_paragraph()
        r2 = p2.add_run(en)
        r2.italic = True
        set_ea(r2)
        for ci in range(1, 8):
            row[ci].text = "○"
            row[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # column widths
    widths = [Cm(9.5)] + [Cm(1.0)] * 7
    for row in table.rows:
        for ci, w in enumerate(widths):
            row.cells[ci].width = w
    doc.add_paragraph()
    return table


# ============================================================
# Title
# ============================================================
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("绿色广告消费者反应问卷（双语版）\nConsumer Responses to Environmental Advertising — Questionnaire (Bilingual)")
r.bold = True
r.font.size = Pt(15)
set_ea(r)

note("【设计说明】本问卷为刺激物问卷设计（stimulus-based survey）。问卷星实施要点见文末附录，附录页不随问卷发放。")
doc.add_paragraph()

# ============================================================
# A. Participant Information & Consent
# ============================================================
h1("A. 研究说明与知情同意 / Participant Information & Consent（首页）")

para("论文题目：环保广告中的感知视觉真实性与漂绿：基于问卷调查的消费者反应研究", bold=True)
para("尊敬的参与者：")
para("感谢您参与本次问卷调查。本调查与我在宁波诺丁汉大学攻读硕士学位的毕业论文相关，研究主题是消费者对环保主题广告的感知与反应。问卷约需 5–8 分钟完成。")
para("您的参与完全自愿。您可以在任何时候退出本次调查，并要求您提供的信息不被用于本研究。您提供的所有信息都将严格保密，仅用于学术研究；在使用您提供的信息时不会涉及您的身份及个人信息。")
para("宁波诺丁汉大学已根据研究伦理审查程序对本研究项目进行了审查，该程序受学校《研究行为与研究伦理守则》约束。如您现在或将来有任何疑问，请联系本人或我的导师。如您对本人的研究行为或研究伦理有任何质疑，请联系我的导师或宁波诺丁汉大学伦理委员会。")
para("研究者：[您的姓名]，[您的UNNC邮箱]；导师：[导师姓名]，[导师邮箱]", italic=True)
doc.add_paragraph()

para("Participant Information Sheet", bold=True)
para("Topic: Perceived Visual Authenticity and Greenwashing in Environmental Advertising: A Survey-Based Study of Consumer Responses", bold=True)
para("Dear Participant,")
para("Thank you for agreeing to participate in this questionnaire survey in connection with my Master's degree at the University of Nottingham Ningbo China. The project studies consumers' perceptions of and responses to environmental-themed advertising. The survey takes about 5–8 minutes.")
para("Your participation in the survey is voluntary. You are able to withdraw from the survey at any time and to request that the information you have provided is not used in the project. Any information provided will be confidential and used for academic purposes only. Your identity will not be disclosed in any use of the information you have supplied during the survey.")
para("The research project has been reviewed according to the ethical review processes in place at the University of Nottingham Ningbo China. These processes are governed by the University's Code of Research Conduct and Research Ethics. Should you have any question now or in the future, please contact me or my supervisor. Should you have concerns related to my conduct of the survey or research ethics, please contact my supervisor or the University's Ethics Committee.")
para("Researcher: [Name], [UNNC email]; Supervisor: [Name], [email]", italic=True)
para("Yours truly,")
doc.add_paragraph()

para("A1 您是否已阅读以上说明，并同意自愿参与？\nHave you read the above information and do you agree to participate voluntarily?（单选）*", bold=True)
options(["是 / Yes", "否（结束问卷）/ No (end survey)"])
doc.add_paragraph()

# ============================================================
# B. Screening & concept explanation
# ============================================================
h1("B. 筛选与概念说明 / Screening & Concept Explanation")

para("B1 您是否年满 18 周岁？ Are you 18 years of age or older?（单选）*", bold=True)
options(["是 / Yes", "否（结束问卷）/ No (end survey)"])
doc.add_paragraph()

para("B2 您是否知道 Nike（耐克）或 Adidas（阿迪达斯）这两个运动品牌？（听说过/浏览过/购买过都算）\nAre you familiar with the sportswear brands Nike or Adidas (heard of / browsed / purchased)?（单选）*", bold=True)
options(["是 / Yes", "否（结束问卷）/ No (end survey)"])
doc.add_paragraph()

para("B3 概念说明 / Concept Explanation", bold=True)
para("中文：本问卷中的“环保主题广告（绿色广告）”指品牌在广告中通过自然景观、绿色色调、回收材料等画面元素或文字，传达环保、可持续等理念的广告。稍后您将看到一则此类广告，请仔细观看后回答问题。")
para("English: In this survey, an “environmental-themed (green) advertisement” refers to an advertisement in which a brand communicates ideas such as environmental protection and sustainability through visual elements (e.g., natural landscapes, green tones, recycled materials) or text. You will shortly view one such advertisement; please look at it carefully before answering the questions.")
doc.add_paragraph()

# ============================================================
# C. Controls: consumption & brand experience (pre-stimulus)
# ============================================================
h1("C. 运动服饰消费与品牌经历（控制变量）/ Sportswear Consumption & Brand Experience (Controls)")

para("C1 过去 12 个月，您大约购买过几次运动服饰（鞋、服装、配件）？\nIn the past 12 months, how many times did you purchase sportswear (shoes, apparel, accessories)?（单选）*", bold=True)
options(["0次 / None", "1–2次 / 1–2 times", "3–5次 / 3–5 times", "6次及以上 / 6 times or more"])
doc.add_paragraph()

para("C2 您对以下品牌的熟悉程度？ How familiar are you with the following brands?（矩阵题；1=非常不熟悉 Very unfamiliar，7=非常熟悉 Very familiar）*", bold=True)
matrix([
    ("C2a", "Nike（耐克）", "Nike"),
    ("C2b", "Adidas（阿迪达斯）", "Adidas"),
], header="品牌 / Brand")

para("C3 过去 12 个月，您是否购买过以下品牌的产品？（可多选）\nIn the past 12 months, have you purchased products from the following brands? (Multiple choice)*", bold=True)
checkboxes(["Nike / 耐克", "Adidas / 阿迪达斯", "以上都没有 / None of the above"])
doc.add_paragraph()

# ============================================================
# D. Stimulus
# ============================================================
h1("D. 广告观看 / Advertisement Viewing")

para("中文：请您仔细观看下面这则广告至少 15 秒，然后根据您的真实感受回答后面的问题。后续问题中“这则广告”指下方广告，“该品牌”指广告中出现的品牌。", bold=True)
para("English: Please look at the advertisement below carefully for at least 15 seconds, and then answer the following questions based on your genuine impressions. In the questions that follow, “this advertisement” refers to the ad below, and “this brand” refers to the brand shown in it.", bold=True)
doc.add_paragraph()
ph = doc.add_paragraph()
ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ph.add_run("【在此插入环保主题广告图片 / Insert environmental-themed advertisement image here】")
r.bold = True
r.font.size = Pt(13)
set_ea(r)
doc.add_paragraph()
note("【设计说明】准备 A、B 两个版本：版本A = 改编自 Nike 环保主题广告；版本B = 改编自 Adidas 环保主题广告。两版除品牌标识外画面结构尽量一致；在问卷星中用“随机展示”将受访者随机分配到其中一版，后续所有题目两版完全相同。建议设置本页最短停留时间 15 秒。")

para("D1 您之前是否见过这则（或类似的）广告？\nHave you seen this (or a similar) advertisement before?（单选）*", bold=True)
options(["见过 / Yes", "没见过 / No", "不确定 / Not sure"])
doc.add_paragraph()

# ============================================================
# E. Main scales
# ============================================================
h1("E. 核心量表（7点同意度）/ Main Scales (7-point agreement)")

para("统一说明 / Common instruction", bold=True)
para("中文：请根据您看完广告后的真实感受，选择您对下列陈述的同意程度。")
para("English: Based on your genuine impressions after viewing the advertisement, please indicate how much you agree with each statement.")
para("量表 / Scale：1=非常不同意 Strongly disagree；2=不同意 Disagree；3=比较不同意 Somewhat disagree；4=一般 Neutral；5=比较同意 Somewhat agree；6=同意 Agree；7=非常同意 Strongly agree")
doc.add_paragraph()

h2("E1 感知视觉真实性 PVA / Perceived Visual Authenticity（矩阵题；6题）")
note("【来源】自编量表，基于 Morhart et al. (2015)、Napoli et al. (2014)、Grayson & Martinec (2004) 改编，对应论文2.3节的维度化方案。")
matrix([
    ("PVA1", "这则广告的画面给我的感觉是真诚的。", "The visuals of this advertisement come across as sincere to me."),
    ("PVA2", "我觉得这则广告的视觉呈现（画面、色彩、场景）是可信的。", "I find the visual presentation of this advertisement (imagery, colours, settings) credible."),
    ("PVA3", "这则广告的画面更像真实的记录，而不是刻意摆拍或包装出来的。", "The visuals of this advertisement feel like a genuine record rather than something staged or packaged."),
    ("PVA4", "我认为画面中的自然或环保元素真实反映了该品牌的环保承诺。", "I think the natural or green elements in the visuals genuinely reflect this brand's environmental commitment."),
    ("PVA5", "这则广告的视觉风格与我对该品牌实际环保表现的印象是一致的。", "The visual style of this advertisement is consistent with my impression of this brand's actual environmental performance."),
    ("PVA6", "这则广告画面的“绿色”程度恰如其分，没有刻意渲染。", "The “green” look of this advertisement feels proportionate rather than overdone."),
])

h2("E2 漂绿感知 PGW / Perceived Greenwashing（矩阵题；5题）")
note("【来源】改编自 Chen & Chang (2013) greenwash 量表，由产品层面调整为广告/品牌层面。")
matrix([
    ("PGW1", "这则广告在环保方面通过画面或图像误导消费者。", "This advertisement misleads consumers with its visuals or graphics regarding the brand's environmental practices."),
    ("PGW2", "这则广告传达的环保信息含糊其辞、难以验证。", "The green message conveyed by this advertisement is vague or difficult to verify."),
    ("PGW3", "这则广告夸大了该品牌实际的环保表现。", "This advertisement overstates or exaggerates the brand's actual environmental performance."),
    ("PGW4", "这则广告掩盖或省略了重要信息，使该品牌显得比实际更环保。", "This advertisement leaves out or masks important information, making the brand seem greener than it really is."),
    ("PGW5", "总体而言，我认为这则广告是在“漂绿”（用绿色形象包装品牌，而非反映真正的环保行动）。", "Overall, I think this advertisement is greenwashing (dressing the brand in a green image rather than reflecting genuine environmental action)."),
])

h2("E3 注意力检查 / Attention Check（1题）")
matrix([
    ("ATT1", "为保证数据质量，请在本题选择“同意（6）”。", "To ensure data quality, please select “Agree (6)” for this item."),
])

h2("E4 绿色信任 GT / Green Trust（矩阵题；5题）")
note("【来源】改编自 Chen (2010) green trust 量表，由产品层面调整为品牌层面。")
matrix([
    ("GT1", "我觉得该品牌的环保声誉总体上是可靠的。", "I feel that this brand's environmental reputation is generally reliable."),
    ("GT2", "我觉得该品牌的环保表现总体上是可信赖的。", "I feel that this brand's environmental performance is generally dependable."),
    ("GT3", "我觉得该品牌的环保主张总体上是值得信任的。", "I feel that this brand's environmental claims are generally trustworthy."),
    ("GT4", "该品牌对环境的关注程度符合我的期望。", "This brand's environmental concern meets my expectations."),
    ("GT5", "我相信该品牌会信守它在环保方面的承诺。", "I believe that this brand keeps its promises and commitments for environmental protection."),
])

h2("E5 购买意愿 PI / Purchase Intention（矩阵题；4题）")
note("【来源】改编自 Chen & Chang (2012) green purchase intention 量表；PI4 为自编补充题。")
matrix([
    ("PI1", "因为该品牌对环境的关注，我愿意购买它的产品。", "I intend to purchase products from this brand because of its environmental concern."),
    ("PI2", "因为该品牌的环保表现，我未来会考虑购买它的产品。", "I expect to purchase products from this brand in the future because of its environmental performance."),
    ("PI3", "总体而言，我乐于购买该品牌的产品，因为它对环境友好。", "Overall, I am glad to purchase products from this brand because they are environmentally friendly."),
    ("PI4", "下次购买运动服饰时，我会优先考虑该品牌。", "I would give priority to this brand the next time I buy sportswear."),
])

# ============================================================
# F. Dispositional controls (post-DV)
# ============================================================
h1("F. 个人环保态度（控制变量）/ Personal Green Dispositions (Controls)")
note("【设计说明】环境关注与广告怀疑倾向放在核心量表之后测量，避免在观看广告前启动（prime）受访者的环保心态，污染核心变量。")

h2("F1 环境关注 EC / Environmental Concern（矩阵题；3题）")
note("【来源】改编自 Kim & Choi (2005)。")
matrix([
    ("EC1", "我非常关心环境问题。", "I am very concerned about environmental issues."),
    ("EC2", "为了保护环境，我愿意改变自己的消费习惯。", "I am willing to change my consumption habits to protect the environment."),
    ("EC3", "环境保护对我个人而言是重要的事。", "Environmental protection is personally important to me."),
])

h2("F2 绿色广告怀疑倾向 GAS / Green Advertising Skepticism（矩阵题；3题）")
note("【来源】改编自 Mohr, Eroğlu & Ellen (1998) 与 Obermiller & Spangenberg (1998)。")
matrix([
    ("GAS1", "大多数广告中的环保说法是夸大其词的。", "Most environmental claims made in advertising are exaggerated."),
    ("GAS2", "广告里的环保宣传多半是为了卖货，而不是真的为了环境。", "Green appeals in advertising are mostly intended to sell products rather than to help the environment."),
    ("GAS3", "我通常不太相信广告中的环保宣传。", "I generally do not trust the environmental claims made in advertisements."),
])

# ============================================================
# G. Demographics
# ============================================================
h1("G. 人口统计信息 / Demographics")

para("G1 性别 / Gender（单选）*", bold=True)
options(["男 / Male", "女 / Female", "其他/不便透露 / Other / Prefer not to say"])
doc.add_paragraph()

para("G2 年龄 / Age（单选）*", bold=True)
options(["18–22", "23–27", "28–35", "36–45", "46–55", "56+"])
doc.add_paragraph()

para("G3 最高学历 / Highest education（单选）*", bold=True)
options(["高中及以下 / High school or below", "大专 / Associate/College", "本科 / Bachelor", "硕士及以上 / Master or above"])
doc.add_paragraph()

para("G4 个人月可支配收入（或生活费/自由支配预算）/ Monthly disposable income（单选）*", bold=True)
options(["3000元以下 / Below 3,000 RMB", "3000–5999元 / 3,000–5,999", "6000–9999元 / 6,000–9,999",
         "10000–19999元 / 10,000–19,999", "20000元及以上 / 20,000+", "不便透露 / Prefer not to say"])
doc.add_paragraph()

para("G5 您目前的身份 / Current occupation status（单选）*", bold=True)
options(["全日制学生 / Full-time student", "在职 / Employed", "自由职业 / Self-employed or freelance", "其他 / Other"])
doc.add_paragraph()

# ============================================================
# H. Open-ended
# ============================================================
h1("H. 开放问题（选答）/ Open-Ended Questions (Optional)")
para("中文：以下问题为选答题，可直接跳过。问题没有标准答案，用一两句话写下您的真实想法即可。您的回答仅用于学术研究，与前面的选择题一样匿名处理。")
para("English: The following questions are optional and may be skipped. There are no right or wrong answers; one or two sentences in your own words are enough. Your responses are anonymous and will be used for academic research only.")
doc.add_paragraph()

para("H1 这则广告的画面中，哪些地方让您觉得“真实”或“不真实”？请简单说明。\nWhich aspects of this advertisement's visuals felt “authentic” or “inauthentic” to you? Please briefly explain.", bold=True)
blank_line()
doc.add_paragraph()

para("H2 一般来说，什么样的画面或视觉元素会让您觉得一个品牌的环保广告可信？\nIn general, what kind of imagery or visual elements make a brand's environmental advertising believable to you?", bold=True)
blank_line()
doc.add_paragraph()

para("H3 您如何看待运动品牌的环保宣传？\nHow do you view environmental claims made by sportswear brands in general?", bold=True)
blank_line()
doc.add_paragraph()

# ============================================================
# I. Closing
# ============================================================
h1("I. 结束语 / Closing")
para("中文：感谢您的参与！如您对问卷有任何疑问，可通过问卷首页所留联系方式联系研究者。")
para("English: Thank you for your participation! If you have any questions, please contact the researcher via the information provided on the first page.")

# ============================================================
# Appendix (not distributed)
# ============================================================
doc.add_page_break()
h1("附录：量表来源与实施说明（本页不随问卷发放，供论文方法章使用）")

para("一、构念—量表来源对照表（可直接改写为论文 Table 3.1）", bold=True)
tbl = doc.add_table(rows=8, cols=5)
tbl.style = "Table Grid"
heads = ["构念 Construct", "编码", "题数", "来源 Source", "对应假设"]
rows = [
    ["感知视觉真实性 Perceived Visual Authenticity", "PVA1–6", "6", "自编；基于 Morhart et al. (2015); Napoli et al. (2014); Grayson & Martinec (2004)", "H1, H2 自变量"],
    ["漂绿感知 Perceived Greenwashing", "PGW1–5", "5", "Chen & Chang (2013)，由产品层面改编至广告/品牌层面", "H1 因变量；H3b 自变量"],
    ["绿色信任 Green Trust", "GT1–5", "5", "Chen (2010)，由产品层面改编至品牌层面", "H2 因变量；H3a 自变量"],
    ["购买意愿 Purchase Intention", "PI1–4", "4", "Chen & Chang (2012)；PI4 自编", "H3 因变量"],
    ["环境关注 Environmental Concern", "EC1–3", "3", "Kim & Choi (2005)", "控制变量"],
    ["绿色广告怀疑 Green Ad Skepticism", "GAS1–3", "3", "Mohr, Eroğlu & Ellen (1998); Obermiller & Spangenberg (1998)", "控制变量"],
    ["品牌熟悉度/购买经历/见过广告", "C2, C3, D1", "4", "自编（事实题）", "控制变量"],
]
for ci, htext in enumerate(heads):
    cell = tbl.rows[0].cells[ci]
    cell.text = htext
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True
            set_ea(r)
for ri, row in enumerate(rows, start=1):
    for ci, val in enumerate(row):
        tbl.rows[ri].cells[ci].text = val
doc.add_paragraph()

para("二、问卷星实施要点", bold=True)
for line in [
    "1. 刺激物随机：用问卷星“随机展示/分组”功能将受访者随机分配到版本A（Nike）或版本B（Adidas），两版仅广告图不同，题目完全相同；后台记录版本变量用于分析。",
    "2. 广告页设置最短停留时间（建议15秒），矩阵题全部设为必答；A1、B1、B2 选“否”自动跳转至结束页。",
    "3. 注意力检查 ATT1 未选“6”的样本在数据清洗阶段剔除（论文4.1节报告剔除数量）。",
    "4. 题目顺序逻辑：品牌熟悉度等事实题放在看广告前；环境关注（EC）与广告怀疑（GAS）放在核心量表之后，避免提前启动环保心态。",
    "5. 中英对照同屏呈现（如示例问卷），如目标样本以中文使用者为主，亦可只发中文版、英文版留作论文附录；正式发放前先做 10–15 人预调研检查题项理解与作答时长。",
    "6. 量表均为 7 点 Likert；PVA 为自编量表，正式分析前建议先做信度（Cronbach's α）与探索性因子分析（EFA）。",
]:
    para(line)

doc.save("绿色广告视觉真实性问卷_双语版.docx")
print("saved")
